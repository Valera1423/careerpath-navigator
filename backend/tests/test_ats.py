def test_ats_analyze(client, auth_headers):
    client.post(
        "/api/v1/users/onboarding",
        json={
            "max_user_id": "test-user-1",
            "desired_position": "Аналитик данных",
            "experience": "none",
            "skills": [],
        },
    )
    resume_text = b"Опыт: Python, SQL. Навыки: pandas. Образование: ВШЭ. Контакты: test@test.ru"
    vacancy_text = "Требуется аналитик данных со знанием SQL, Python, статистики и Power BI"

    files = {"file": ("resume.txt", resume_text, "text/plain")}
    # TXT не поддерживается — используем docx-набор через отдельный тест
    r = client.post(
        "/api/v1/ats/analyze",
        files=files,
        data={"vacancy_text": vacancy_text},
        headers=auth_headers,
    )
    assert r.status_code == 400  # txt не поддерживается


def test_ats_analyze_docx(client, auth_headers):
    import io
    from docx import Document

    client.post(
        "/api/v1/users/onboarding",
        json={
            "max_user_id": "test-user-1",
            "desired_position": "Аналитик данных",
            "experience": "none",
            "skills": [],
        },
    )

    doc = Document()
    doc.add_paragraph("Опыт: Python, SQL")
    doc.add_paragraph("Навыки: pandas, statistics")
    doc.add_paragraph("Образование: ВШЭ")
    doc.add_paragraph("Контакты: test@test.ru")
    buf = io.BytesIO()
    doc.save(buf)

    files = {
        "file": (
            "resume.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    r = client.post(
        "/api/v1/ats/analyze",
        files=files,
        data={"vacancy_text": "Требуется аналитик данных: SQL, Python, statistics"},
        headers=auth_headers,
    )
    assert r.status_code == 200
    body = r.json()
    assert body["score"] > 0
    assert "sql" in body["matched_keywords"]